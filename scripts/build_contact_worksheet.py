#!/usr/bin/env python3
"""Attach known businesses to an address list and lay out a contact-research worksheet.

Takes a user-supplied address/PID list (e.g. a low-site-coverage shortlist) and fills in
what municipal open data can actually tell us — which businesses hold a licence at each
address — then leaves clearly-labelled empty columns for contact details that must be
sourced elsewhere.

    python3 scripts/build_contact_worksheet.py addresses.xlsx \
        --delta-business data/samples/municipal/delta_business_licences.csv \
        --delta-parcels data/samples/municipal/delta_parcels.csv \
        -o output/low_coverage_contacts.xlsx

Deliberately does NOT invent contact names, emails, or phone numbers. Every populated cell
traces to a named source in the Source column; anything not sourced is left blank for
research rather than filled with a plausible guess. A broker acting on a fabricated contact
is worse off than one acting on a blank cell.
"""

import argparse
import csv
import re
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

FONT = "Arial"

STREET_ABBREV = [
    ("AVENUE", "AVE"), ("STREET", "ST"), ("BOULEVARD", "BLVD"), ("PLACE", "PL"),
    ("ROAD", "RD"), ("DRIVE", "DR"), ("CRESCENT", "CR"), ("PARKWAY", "PKWY"),
    ("HIGHWAY", "HWY"), ("LANE", "LN"), ("COURT", "CRT"),
]


def norm_address(a) -> str:
    """Normalize for joining. Delta's licence ADDRESS field embeds a newline plus
    'DELTA BC V4G 1A7', and prefixes units as '102-', 'E-', 'DOCK-' — both have to go
    before a street address will match a clean list."""
    if not a:
        return ""
    a = str(a).split("\n")[0].upper().strip()
    a = re.sub(r"^[A-Z0-9]{1,4}\s*-\s*", "", a)
    a = re.sub(r"^UNIT\s+\S+,?\s*", "", a)
    a = re.sub(r"^#\S+,?\s*", "", a)
    a = re.sub(r"\s+", " ", a)
    for long, short in STREET_ABBREV:
        a = re.sub(r"\b" + long + r"\b", short, a)
    return a.strip()


def load_business_index(path: Path) -> dict:
    index = {}
    with open(path, encoding="utf-8-sig") as f:
        for row in csv.DictReader(f):
            index.setdefault(norm_address(row["ADDRESS"]), []).append(row)
    return index


def load_pid_set(path: Path) -> set:
    pids = set()
    with open(path, encoding="utf-8-sig") as f:
        for row in csv.DictReader(f):
            if row.get("PID"):
                pids.add(row["PID"].strip())
    return pids


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("addresses_xlsx", type=Path)
    ap.add_argument("--delta-business", type=Path, required=True)
    ap.add_argument("--delta-parcels", type=Path, required=True)
    ap.add_argument("-o", "--out", type=Path, required=True)
    args = ap.parse_args()

    biz = load_business_index(args.delta_business)
    delta_pids = load_pid_set(args.delta_parcels)

    wb_in = openpyxl.load_workbook(args.addresses_xlsx, data_only=True)
    ws_in = wb_in.active
    src_rows = [r for r in ws_in.iter_rows(min_row=2, values_only=True) if any(r)]

    out = []
    for order, pid, address, coverage in src_rows:
        pid = str(pid).strip() if pid else ""
        municipality = "Delta" if pid in delta_pids else "Richmond (assumed)"
        matches = biz.get(norm_address(address), []) if municipality == "Delta" else []

        names = [m["TRADE_NAME"].strip() for m in matches]
        types = sorted({m["BUSINESS_TYPE"].split("/")[0].strip() for m in matches if m.get("BUSINESS_TYPE")})
        employees = [m["EMPLOYEES"] for m in matches if m.get("EMPLOYEES")]

        if matches:
            source = "City of Delta business licence open data"
        elif municipality == "Delta":
            source = "No Delta licence at this address — needs research"
        else:
            source = "Richmond publishes no business licence data — needs research"

        out.append({
            "order": order,
            "pid": pid,
            "address": str(address).strip() if address else "",
            "coverage": coverage,
            "municipality": municipality,
            "businesses": "; ".join(names),
            "business_count": len(names),
            "business_type": "; ".join(types),
            "employees": "; ".join(str(e) for e in employees),
            "source": source,
        })

    # --- write workbook ---
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Contact Worksheet"

    filled_headers = ["Order", "PID", "Address", "Site Coverage %", "Municipality",
                      "Company / Companies at Address", "# Businesses", "Business Type",
                      "Employees", "Source (for company column)"]
    research_headers = ["Main Business Phone", "Company Website",
                        "Senior Contact 1 - Name", "Contact 1 - Title", "Contact 1 - Email", "Contact 1 - Phone",
                        "Senior Contact 2 - Name", "Contact 2 - Title", "Contact 2 - Email", "Contact 2 - Phone",
                        "Contact Source / Verified?"]
    ws.append(filled_headers + research_headers)

    for r in out:
        ws.append([
            r["order"], r["pid"], r["address"], r["coverage"], r["municipality"],
            r["businesses"] or "", r["business_count"] or "", r["business_type"],
            r["employees"], r["source"],
        ] + [""] * len(research_headers))

    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    filled_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    research_fill = PatternFill(start_color="7F6000", end_color="7F6000", fill_type="solid")

    for idx, cell in enumerate(ws[1], start=1):
        cell.font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
        cell.fill = filled_fill if idx <= len(filled_headers) else research_fill
        cell.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
        cell.border = border

    empty_research = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
    no_data = PatternFill(start_color="F8CBCB", end_color="F8CBCB", fill_type="solid")

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            cell.font = Font(name=FONT, size=10)
            cell.border = border
            cell.alignment = Alignment(vertical="top",
                                       wrap_text=(cell.column_letter in ("C", "F", "H", "J")))
        # company column empty -> flag red so gaps are obvious, never silently blank
        if not row[5].value:
            row[5].fill = no_data
        for i in range(len(filled_headers), len(filled_headers) + len(research_headers)):
            row[i].fill = empty_research

    widths = [6, 13, 22, 13, 16, 34, 11, 24, 11, 30,
              16, 22, 20, 18, 24, 16, 20, 18, 24, 16, 24]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = ws.dimensions

    # --- notes ---
    delta_rows = [r for r in out if r["municipality"] == "Delta"]
    delta_matched = [r for r in delta_rows if r["business_count"]]
    rich_rows = [r for r in out if r["municipality"] != "Delta"]

    notes = wb.create_sheet("Notes")
    notes["A1"] = "Low Site Coverage List — Company & Contact Worksheet"
    notes["A1"].font = Font(name=FONT, bold=True, size=13)
    lines = [
        "",
        "BLUE columns are filled from official data and are traceable via the Source column.",
        "GOLD columns are intentionally blank — they need research and must not be guessed.",
        "",
        f"Total addresses: {len(out)}",
        f"  Delta: {len(delta_rows)} — {len(delta_matched)} matched to a business licence "
        f"({100*len(delta_matched)/len(delta_rows):.0f}%)",
        f"  Richmond: {len(rich_rows)} — 0 matched (see below)",
        "",
        "WHY RICHMOND IS EMPTY:",
        "Richmond is the one Metro Vancouver municipality checked that publishes no open data",
        "portal at all — no business licence dataset, no parcel API, only a read-only map viewer.",
        "Surrey, Delta, Burnaby, Coquitlam and Langley all publish machine-readable data;",
        "Richmond does not. Those addresses need per-address lookup, not a bulk join.",
        "",
        "WHY THE CONTACT COLUMNS ARE BLANK — and what will actually fill them:",
        "Delta's business licence data contains company name, business type and employee count.",
        "It contains no phone, no email, and no personal names. Tested against live web search:",
        "  - A main business phone IS findable per company (verified on two of these companies).",
        "  - Named senior contacts and their email addresses are NOT findable this way for",
        "    private industrial companies — general web results return switchboard numbers and",
        "    generic info@ addresses, not named decision-makers.",
        "",
        "The authoritative source for named senior people is a BC Registries corporate search,",
        "which lists directors and officers — the same search type already held for 15 companies",
        "in this project. It is paid and per-company, but it is the only source that gives a real,",
        "verifiable name attached to the entity that actually owns or operates the business.",
        "",
        "NOTHING IN THIS SHEET IS INVENTED. No contact name, email, or phone number has been",
        "generated or pattern-guessed. A plausible-looking wrong contact is worse than a blank",
        "cell — it costs a call, an email to the wrong person, and credibility.",
    ]
    for i, line in enumerate(lines, start=2):
        notes[f"A{i}"] = line
        notes[f"A{i}"].font = Font(name=FONT, size=10)
    notes.column_dimensions["A"].width = 100

    args.out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(args.out)
    print(f"{len(out)} addresses -> {args.out}")
    print(f"  Delta {len(delta_rows)} ({len(delta_matched)} with businesses), Richmond {len(rich_rows)}")


if __name__ == "__main__":
    main()
